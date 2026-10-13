import pdfplumber
import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
import re
import os
from pathlib import Path
import sys
sys.stdout.reconfigure(encoding="utf-8")


def extrair_informacoes_pdf(caminho_pdf):

    nome_empresa = None
    valor_produtos = None
    valor_nota = None
    
    with pdfplumber.open(caminho_pdf) as pdf:
        total_paginas = len(pdf.pages)
        
        # Vou analisar as páginas de trás para frente (última → penúltima → etc.)
        paginas_para_analisar = list(reversed(range(total_paginas)))
        
        texto_completo = ""
        
        for num_pagina in paginas_para_analisar:
            pagina = pdf.pages[num_pagina]
            texto = pagina.extract_text() or ""
            texto_completo += "\n" + texto
            

            if nome_empresa is None:
                # Regex: captura o código (números/números) + o nome que vem depois
                match_empresa = re.search(
                    r'(\d+/\d+)\s+([A-ZÀ-Ü0-9\s\.\-&]+?)(?:\s{2,}|\||$)',
                    texto,
                    re.IGNORECASE
                )
                if match_empresa:
                    nome_empresa = match_empresa.group(2).strip()
                    # Remove espaços extras no final
                    nome_empresa = re.sub(r'\s+', ' ', nome_empresa).strip()
            

            if valor_produtos is None:
                match_produtos = re.search(
                    r'Valor\s+Total\s+dos\s+Produtos\s*=>\s*([\d.,]+)',
                    texto,
                    re.IGNORECASE
                )
                if match_produtos:
                    valor_produtos = match_produtos.group(1).strip()
            
            if valor_nota is None:
                match_nota = re.search(
                    r'Valor\s+Total\s+da\s+Nota\s*=>\s*([\d.,]+)',
                    texto,
                    re.IGNORECASE
                )
                if match_nota:
                    valor_nota = match_nota.group(1).strip()
            
            # Se já encontrou tudo, não preciso continuar analisando as páginas
            if nome_empresa and valor_produtos and valor_nota:
                break
    
    return {
        "nome_empresa": nome_empresa or "NÃO ENCONTRADO",
        "valor_produtos": valor_produtos or "NÃO ENCONTRADO",
        "valor_nota": valor_nota or "NÃO ENCONTRADO"
    }


def criar_ou_atualizar_planilha(dados, nome_arquivo_pdf, caminho_excel="resultado_pedidos.xlsx"):
    """
    Cria a planilha Excel (se não existir) ou adiciona uma nova linha (se já existir).
    """
    
    # Verifico se a planilha já existe
    if os.path.exists(caminho_excel):
        wb = openpyxl.load_workbook(caminho_excel)
        ws = wb.active
    else:
        # Crio uma planilha nova
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Pedidos Extraídos"
        
        # Cabeçalhos
        cabecalhos = [
            "Arquivo PDF",
            "Nome da Empresa",
            "Valor Total dos Produtos",
            "Valor Total da Nota"
        ]
        
        # Estilo do cabeçalho
        fonte_cabecalho = Font(bold=True, color="FFFFFF")
        preenchimento = PatternFill(start_color="2F5496", end_color="2F5496", fill_type="solid")
        alinhamento = Alignment(horizontal="center", vertical="center")
        borda = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        
        for col, titulo in enumerate(cabecalhos, 1):
            celula = ws.cell(row=1, column=col, value=titulo)
            celula.font = fonte_cabecalho
            celula.fill = preenchimento
            celula.alignment = alinhamento
            celula.border = borda
        
        # Ajusta largura das colunas
        ws.column_dimensions['A'].width = 50
        ws.column_dimensions['B'].width = 45
        ws.column_dimensions['C'].width = 25
        ws.column_dimensions['D'].width = 22
    
    # Adiciona a nova linha com os dados
    nova_linha = ws.max_row + 1
    
    ws.cell(row=nova_linha, column=1, value=nome_arquivo_pdf)
    ws.cell(row=nova_linha, column=2, value=dados["nome_empresa"])
    ws.cell(row=nova_linha, column=3, value=dados["valor_produtos"])
    ws.cell(row=nova_linha, column=4, value=dados["valor_nota"])
    
    # Aplica alinhamento e borda nas células
    for col in range(1, 5):
        celula = ws.cell(row=nova_linha, column=col)
        celula.alignment = Alignment(horizontal="center" if col in [2, 3, 4] else "left", vertical="center")
        celula.border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
    
    wb.save(caminho_excel)
    print(f"✅ Dados salvos na planilha: {caminho_excel}")


def processar_pdfs(pasta_ou_arquivo, caminho_excel="resultado_pedidos.xlsx"):

    caminho = Path(pasta_ou_arquivo)
    
    if caminho.is_file() and caminho.suffix.lower() == ".pdf":
        lista_pdfs = [caminho]
    elif caminho.is_dir():
        lista_pdfs = list(caminho.glob("*.pdf"))
    else:
        print("❌ Caminho inválido. Informe um arquivo PDF ou uma pasta.")
        return
    
    if not lista_pdfs:
        print("❌ Nenhum PDF encontrado.")
        return
    
    print(f"🔍 Encontrados {len(lista_pdfs)} PDF(s) para processar...\n")
    
    for pdf in lista_pdfs:
        print(f"📄 Processando: {pdf.name}")
        
        try:
            dados = extrair_informacoes_pdf(str(pdf))
            
            print(f"   → Empresa: {dados['nome_empresa']}")
            print(f"   → Valor Produtos: {dados['valor_produtos']}")
            print(f"   → Valor Nota: {dados['valor_nota']}")
            print()
            
            criar_ou_atualizar_planilha(dados, pdf.name, caminho_excel)
            
        except Exception as e:
            print(f"❌ Erro ao processar {pdf.name}: {e}\n")
    
    print("🎉 Processamento finalizado!")


# COMO USAR O SCRIPT
if __name__ == "__main__":

    processar_pdfs(
        pasta_ou_arquivo= r"C:\\Users\\COBRANÇA\\Desktop\\Giliarde\\python_projetos\\automacao_de_processamento_de_pedidos\\pdfs",   # <-- Coloque aqui o caminho da pasta ou do PDF
        caminho_excel="resultado_pedidos.xlsx"          # <-- Nome da planilha que será gerada
    )
